import os
import docx
from PIL import Image

SAMPLES_DIR = os.path.join(os.path.dirname(__file__), 'samples')

def create_sample_file(filename: str, target_dir: str = None) -> str:
    if target_dir is None:
        target_dir = SAMPLES_DIR
    os.makedirs(target_dir, exist_ok=True)
    target_path = os.path.join(target_dir, filename)

    if filename == "sample_doc.docx":
        doc = docx.Document()
        doc.add_heading("Звіт про цілісність даних", level=1)
        doc.add_paragraph("Цей документ використовується для тестування криптографічної хеш-функції та контролю цілісності.")
        doc.core_properties.author = "Студент"
        doc.core_properties.comments = "Початкова версія документа"
        doc.save(target_path)
    elif filename == "sample_doc.txt":
        with open(target_path, "w", encoding="utf-8") as f:
            f.write("Офіційний текстовий документ для перевірки хеш-суми та пошуку колізій.\n")
    elif filename == "sample_code.py":
        with open(target_path, "w", encoding="utf-8") as f:
            f.write("def calculate_sum(a, b):\n    return a + b\n\nif __name__ == '__main__':\n    print('Сума:', calculate_sum(10, 25))\n")
    elif filename == "sample_image.bmp":
        img = Image.new("RGB", (64, 64), color=(30, 144, 255))
        img.save(target_path, format="BMP")
    return target_path

def ensure_all_samples(target_dir: str = None):
    if target_dir is None:
        target_dir = SAMPLES_DIR
    files = ["sample_doc.docx", "sample_doc.txt", "sample_code.py", "sample_image.bmp"]
    for f in files:
        path = os.path.join(target_dir, f)
        if not os.path.exists(path):
            create_sample_file(f, target_dir)

if __name__ == '__main__':
    ensure_all_samples()
    print("Samples verified in:", SAMPLES_DIR)
